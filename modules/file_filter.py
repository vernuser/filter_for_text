"""
文件内容过滤模块
"""
import os
import hashlib
import mimetypes
import logging
import zipfile
import tarfile
import docx
import PyPDF2
import openpyxl
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Union
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from core.filter_engine import FilterEngine
import threading
import time
import sqlite3
from config.settings import DATABASE_PATH

class FileFilter:
    """文件内容过滤器"""
    
    def __init__(self, filter_engine: FilterEngine):
        self.filter_engine = filter_engine
        self.logger = logging.getLogger(__name__)
        self.db_path = DATABASE_PATH
        self.monitoring = False
        self.observer = None
        self.monitored_paths = set()
        
        # 支持的文件类型
        self.supported_text_extensions = {
            '.txt', '.md', '.py', '.js', '.html', '.css', '.xml', '.json',
            '.csv', '.log', '.ini', '.cfg', '.conf', '.yaml', '.yml'
        }
        
        self.supported_office_extensions = {
            '.docx', '.xlsx', '.pptx', '.pdf'
        }
        
        self.supported_archive_extensions = {
            '.zip', '.rar', '.7z', '.tar', '.gz', '.bz2'
        }
        
        # 危险文件类型
        self.dangerous_extensions = {
            '.exe', '.bat', '.cmd', '.com', '.scr', '.pif', '.vbs', '.js',
            '.jar', '.msi', '.dll', '.sys', '.drv'
        }
        
        self._init_file_database()
    
    def _init_file_database(self):
        """初始化文件过滤数据库"""
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS file_scan_results (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                file_path TEXT NOT NULL,
                file_hash TEXT,
                file_size INTEGER,
                file_type TEXT,
                scan_result TEXT,
                violations_count INTEGER DEFAULT 0,
                risk_level INTEGER DEFAULT 0,
                scan_time TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                action_taken TEXT
            )
        ''')
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS quarantine_files (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                original_path TEXT NOT NULL,
                quarantine_path TEXT NOT NULL,
                file_hash TEXT,
                quarantine_reason TEXT,
                quarantine_time TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        conn.commit()
        conn.close()
    
    def scan_file(self, file_path: str) -> Dict:
        """
        扫描单个文件
        
        Args:
            file_path: 文件路径
            
        Returns:
            Dict: 扫描结果
        """
        try:
            if not os.path.exists(file_path):
                return {'status': 'error', 'message': '文件不存在'}
            
            file_info = self._get_file_info(file_path)
            
            # 检查文件类型安全性
            type_check = self._check_file_type_safety(file_path)
            
            # 扫描文件内容
            content_result = self._scan_file_content(file_path)
            
            # 计算风险级别
            risk_level = self._calculate_risk_level(type_check, content_result)
            
            # 决定处理动作
            action = self._determine_action(risk_level, content_result)
            
            scan_result = {
                'file_path': file_path,
                'file_info': file_info,
                'type_check': type_check,
                'content_result': content_result,
                'risk_level': risk_level,
                'action': action,
                'scan_time': time.time()
            }
            
            # 记录扫描结果
            self._log_scan_result(scan_result)
            
            # 执行处理动作
            if action['type'] in ['quarantine', 'delete']:
                self._execute_action(file_path, action)
            
            return scan_result
            
        except Exception as e:
            self.logger.error(f"文件扫描错误 {file_path}: {e}")
            return {'status': 'error', 'message': str(e)}
    
    def _get_file_info(self, file_path: str) -> Dict:
        """获取文件基本信息"""
        try:
            stat = os.stat(file_path)
            file_hash = self._calculate_file_hash(file_path)
            mime_type, _ = mimetypes.guess_type(file_path)
            
            return {
                'size': stat.st_size,
                'modified_time': stat.st_mtime,
                'hash': file_hash,
                'mime_type': mime_type,
                'extension': Path(file_path).suffix.lower()
            }
        except Exception as e:
            self.logger.error(f"获取文件信息错误: {e}")
            return {}
    
    def _calculate_file_hash(self, file_path: str) -> str:
        """计算文件哈希值"""
        try:
            hash_md5 = hashlib.md5()
            with open(file_path, "rb") as f:
                for chunk in iter(lambda: f.read(4096), b""):
                    hash_md5.update(chunk)
            return hash_md5.hexdigest()
        except Exception as e:
            self.logger.error(f"计算文件哈希错误: {e}")
            return ""
    
    def _check_file_type_safety(self, file_path: str) -> Dict:
        """检查文件类型安全性"""
        extension = Path(file_path).suffix.lower()
        
        if extension in self.dangerous_extensions:
            return {
                'safe': False,
                'reason': 'dangerous_extension',
                'severity': 3,
                'description': f'危险文件类型: {extension}'
            }
        
        # 检查文件头部魔数
        magic_check = self._check_file_magic(file_path)
        if not magic_check['valid']:
            return {
                'safe': False,
                'reason': 'invalid_magic',
                'severity': 2,
                'description': '文件头部不匹配扩展名'
            }
        
        return {
            'safe': True,
            'reason': 'type_safe',
            'severity': 0,
            'description': '文件类型安全'
        }
    
    def _check_file_magic(self, file_path: str) -> Dict:
        """检查文件魔数"""
        try:
            with open(file_path, 'rb') as f:
                header = f.read(16)
            
            # 常见文件类型的魔数
            magic_numbers = {
                b'\x50\x4B\x03\x04': '.zip',
                b'\x50\x4B\x05\x06': '.zip',
                b'\x50\x4B\x07\x08': '.zip',
                b'\x4D\x5A': '.exe',
                b'\x7F\x45\x4C\x46': '.elf',
                b'\x25\x50\x44\x46': '.pdf',
                b'\xFF\xD8\xFF': '.jpg',
                b'\x89\x50\x4E\x47': '.png',
                b'\x47\x49\x46\x38': '.gif',
            }
            
            extension = Path(file_path).suffix.lower()
            
            for magic, expected_ext in magic_numbers.items():
                if header.startswith(magic):
                    return {
                        'valid': extension == expected_ext,
                        'detected_type': expected_ext,
                        'declared_type': extension
                    }
            
            return {'valid': True, 'detected_type': 'unknown', 'declared_type': extension}
            
        except Exception as e:
            self.logger.error(f"检查文件魔数错误: {e}")
            return {'valid': True, 'detected_type': 'unknown', 'declared_type': 'unknown'}
    
    def _scan_file_content(self, file_path: str) -> Dict:
        """扫描文件内容"""
        extension = Path(file_path).suffix.lower()
        
        try:
            if extension in self.supported_text_extensions:
                return self._scan_text_file(file_path)
            elif extension in self.supported_office_extensions:
                return self._scan_office_file(file_path)
            elif extension in self.supported_archive_extensions:
                return self._scan_archive_file(file_path)
            else:
                return self._scan_binary_file(file_path)
                
        except Exception as e:
            self.logger.error(f"扫描文件内容错误 {file_path}: {e}")
            return {'status': 'error', 'violations': [], 'message': str(e)}
    
    def _scan_text_file(self, file_path: str) -> Dict:
        """扫描文本文件"""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            
            # 使用过滤引擎检查内容
            filtered_content, violations = self.filter_engine.filter_text(content)
            
            # 检查URL和IP
            url_violations = self._extract_and_check_urls(content)
            ip_violations = self._extract_and_check_ips(content)
            
            all_violations = violations + url_violations + ip_violations
            
            return {
                'status': 'success',
                'content_type': 'text',
                'violations': all_violations,
                'filtered_content': filtered_content,
                'original_size': len(content),
                'filtered_size': len(filtered_content)
            }
            
        except Exception as e:
            self.logger.error(f"扫描文本文件错误: {e}")
            return {'status': 'error', 'violations': [], 'message': str(e)}
    
    def _scan_office_file(self, file_path: str) -> Dict:
        """扫描Office文件"""
        extension = Path(file_path).suffix.lower()
        violations = []
        
        try:
            if extension == '.docx':
                content = self._extract_docx_text(file_path)
            elif extension == '.xlsx':
                content = self._extract_xlsx_text(file_path)
            elif extension == '.pdf':
                content = self._extract_pdf_text(file_path)
            else:
                return {'status': 'unsupported', 'violations': []}
            
            if content:
                filtered_content, text_violations = self.filter_engine.filter_text(content)
                violations.extend(text_violations)
            
            return {
                'status': 'success',
                'content_type': 'office',
                'violations': violations,
                'extracted_text_length': len(content) if content else 0
            }
            
        except Exception as e:
            self.logger.error(f"扫描Office文件错误: {e}")
            return {'status': 'error', 'violations': [], 'message': str(e)}
    
    def _extract_docx_text(self, file_path: str) -> str:
        """提取DOCX文件文本"""
        try:
            doc = docx.Document(file_path)
            text = []
            for paragraph in doc.paragraphs:
                text.append(paragraph.text)
            return '\n'.join(text)
        except Exception as e:
            self.logger.error(f"提取DOCX文本错误: {e}")
            return ""
    
    def _extract_xlsx_text(self, file_path: str) -> str:
        """提取XLSX文件文本"""
        try:
            workbook = openpyxl.load_workbook(file_path)
            text = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value:
                            text.append(str(cell.value))
            return '\n'.join(text)
        except Exception as e:
            self.logger.error(f"提取XLSX文本错误: {e}")
            return ""
    
    def _extract_pdf_text(self, file_path: str) -> str:
        """提取PDF文件文本"""
        try:
            text = []
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text.append(page.extract_text())
            return '\n'.join(text)
        except Exception as e:
            self.logger.error(f"提取PDF文本错误: {e}")
            return ""
    
    def _scan_archive_file(self, file_path: str) -> Dict:
        """扫描压缩文件"""
        violations = []
        
        try:
            if file_path.endswith('.zip'):
                with zipfile.ZipFile(file_path, 'r') as zip_file:
                    for file_info in zip_file.filelist:
                        # 检查文件名
                        filename_violations = self._check_filename_safety(file_info.filename)
                        violations.extend(filename_violations)
                        
                        # 检查文件大小（压缩炸弹检测）
                        if file_info.file_size > 100 * 1024 * 1024:  # 100MB
                            violations.append({
                                'type': 'large_file_in_archive',
                                'filename': file_info.filename,
                                'size': file_info.file_size,
                                'severity': 2
                            })
            
            return {
                'status': 'success',
                'content_type': 'archive',
                'violations': violations
            }
            
        except Exception as e:
            self.logger.error(f"扫描压缩文件错误: {e}")
            return {'status': 'error', 'violations': [], 'message': str(e)}
    
    def _scan_binary_file(self, file_path: str) -> Dict:
        """扫描二进制文件"""
        violations = []
        
        try:
            # 检查文件大小
            file_size = os.path.getsize(file_path)
            if file_size > 500 * 1024 * 1024:  # 500MB
                violations.append({
                    'type': 'large_binary_file',
                    'size': file_size,
                    'severity': 1
                })
            
            # 简单的二进制内容检查
            with open(file_path, 'rb') as f:
                header = f.read(1024)
                
                # 检查可疑字符串
                suspicious_strings = [b'eval', b'exec', b'system', b'shell']
                for sus_str in suspicious_strings:
                    if sus_str in header:
                        violations.append({
                            'type': 'suspicious_binary_content',
                            'pattern': sus_str.decode('utf-8', errors='ignore'),
                            'severity': 2
                        })
            
            return {
                'status': 'success',
                'content_type': 'binary',
                'violations': violations
            }
            
        except Exception as e:
            self.logger.error(f"扫描二进制文件错误: {e}")
            return {'status': 'error', 'violations': [], 'message': str(e)}
    
    def _extract_and_check_urls(self, content: str) -> List[Dict]:
        """提取并检查URL"""
        import re
        violations = []
        
        url_pattern = r'http[s]?://(?:[a-zA-Z]|[0-9]|[$-_@.&+]|[!*\\(\\),]|(?:%[0-9a-fA-F][0-9a-fA-F]))+'
        urls = re.findall(url_pattern, content)
        
        for url in urls:
            is_allowed, violation = self.filter_engine.filter_url(url)
            if not is_allowed:
                violations.append({
                    'type': 'blocked_url_in_file',
                    'url': url,
                    'violation': violation
                })
        
        return violations
    
    def _extract_and_check_ips(self, content: str) -> List[Dict]:
        """提取并检查IP地址"""
        import re
        violations = []
        
        ip_pattern = r'\b(?:[0-9]{1,3}\.){3}[0-9]{1,3}\b'
        ips = re.findall(ip_pattern, content)
        
        for ip in ips:
            is_allowed, violation = self.filter_engine.filter_ip(ip)
            if not is_allowed:
                violations.append({
                    'type': 'blocked_ip_in_file',
                    'ip': ip,
                    'violation': violation
                })
        
        return violations
    
    def _check_filename_safety(self, filename: str) -> List[Dict]:
        """检查文件名安全性"""
        violations = []
        
        # 检查危险扩展名
        extension = Path(filename).suffix.lower()
        if extension in self.dangerous_extensions:
            violations.append({
                'type': 'dangerous_filename',
                'filename': filename,
                'extension': extension,
                'severity': 3
            })
        
        # 检查路径遍历
        if '..' in filename or filename.startswith('/') or ':' in filename:
            violations.append({
                'type': 'path_traversal',
                'filename': filename,
                'severity': 3
            })
        
        return violations
    
    def _calculate_risk_level(self, type_check: Dict, content_result: Dict) -> int:
        """计算风险级别"""
        risk_level = 0
        
        # 文件类型风险
        if not type_check.get('safe', True):
            risk_level += type_check.get('severity', 0)
        
        # 内容风险
        violations = content_result.get('violations', [])
        if violations:
            max_severity = max([v.get('severity', 1) for v in violations])
            risk_level += max_severity
        
        return min(risk_level, 5)  # 最高风险级别为5
    
    def _determine_action(self, risk_level: int, content_result: Dict) -> Dict:
        """确定处理动作"""
        if risk_level >= 4:
            return {'type': 'quarantine', 'reason': 'high_risk'}
        elif risk_level >= 2:
            return {'type': 'warn', 'reason': 'medium_risk'}
        else:
            return {'type': 'allow', 'reason': 'low_risk'}
    
    def _execute_action(self, file_path: str, action: Dict):
        """执行处理动作"""
        try:
            if action['type'] == 'quarantine':
                self._quarantine_file(file_path, action['reason'])
            elif action['type'] == 'delete':
                os.remove(file_path)
                self.logger.info(f"删除危险文件: {file_path}")
        except Exception as e:
            self.logger.error(f"执行动作错误: {e}")
    
    def _quarantine_file(self, file_path: str, reason: str):
        """隔离文件"""
        try:
            quarantine_dir = os.path.join(os.path.dirname(self.db_path), 'quarantine')
            os.makedirs(quarantine_dir, exist_ok=True)
            
            filename = os.path.basename(file_path)
            timestamp = int(time.time())
            quarantine_path = os.path.join(quarantine_dir, f"{timestamp}_{filename}")
            
            # 移动文件到隔离区
            os.rename(file_path, quarantine_path)
            
            # 记录隔离信息
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            cursor.execute('''
                INSERT INTO quarantine_files (original_path, quarantine_path, file_hash, quarantine_reason)
                VALUES (?, ?, ?, ?)
            ''', (file_path, quarantine_path, self._calculate_file_hash(quarantine_path), reason))
            
            conn.commit()
            conn.close()
            
            self.logger.info(f"文件已隔离: {file_path} -> {quarantine_path}")
            
        except Exception as e:
            self.logger.error(f"隔离文件错误: {e}")
    
    def _log_scan_result(self, scan_result: Dict):
        """记录扫描结果"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            violations_count = len(scan_result.get('content_result', {}).get('violations', []))
            
            cursor.execute('''
                INSERT INTO file_scan_results 
                (file_path, file_hash, file_size, file_type, scan_result, violations_count, risk_level, action_taken)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                scan_result['file_path'],
                scan_result.get('file_info', {}).get('hash', ''),
                scan_result.get('file_info', {}).get('size', 0),
                scan_result.get('file_info', {}).get('extension', ''),
                str(scan_result),
                violations_count,
                scan_result.get('risk_level', 0),
                scan_result.get('action', {}).get('type', 'unknown')
            ))
            
            conn.commit()
            conn.close()
            
        except Exception as e:
            self.logger.error(f"记录扫描结果错误: {e}")
    
    def start_monitoring(self, paths: List[str]):
        """开始监控指定路径"""
        try:
            if self.monitoring:
                self.stop_monitoring()
            
            self.observer = Observer()
            event_handler = FileFilterEventHandler(self)
            
            for path in paths:
                if os.path.exists(path):
                    self.observer.schedule(event_handler, path, recursive=True)
                    self.monitored_paths.add(path)
                    self.logger.info(f"开始监控路径: {path}")
            
            self.observer.start()
            self.monitoring = True
            self.logger.info("文件监控已启动")
            
        except Exception as e:
            self.logger.error(f"启动文件监控错误: {e}")
    
    def stop_monitoring(self):
        """停止文件监控"""
        if self.observer:
            self.observer.stop()
            self.observer.join()
            self.monitoring = False
            self.monitored_paths.clear()
            self.logger.info("文件监控已停止")
    
    def get_scan_stats(self) -> Dict:
        """获取扫描统计信息"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            # 总扫描次数
            cursor.execute('SELECT COUNT(*) FROM file_scan_results')
            total_scans = cursor.fetchone()[0]
            
            # 按风险级别统计
            cursor.execute('''
                SELECT risk_level, COUNT(*) 
                FROM file_scan_results 
                GROUP BY risk_level
            ''')
            risk_stats = dict(cursor.fetchall())
            
            # 隔离文件数量
            cursor.execute('SELECT COUNT(*) FROM quarantine_files')
            quarantined_files = cursor.fetchone()[0]
            
            conn.close()
            
            return {
                'total_scans': total_scans,
                'risk_distribution': risk_stats,
                'quarantined_files': quarantined_files,
                'monitoring_active': self.monitoring,
                'monitored_paths': list(self.monitored_paths)
            }
            
        except Exception as e:
            self.logger.error(f"获取扫描统计错误: {e}")
            return {}


class FileFilterEventHandler(FileSystemEventHandler):
    """文件系统事件处理器"""
    
    def __init__(self, file_filter: FileFilter):
        self.file_filter = file_filter
        self.logger = logging.getLogger(__name__)
    
    def on_created(self, event):
        """文件创建事件"""
        if not event.is_directory:
            self.logger.info(f"检测到新文件: {event.src_path}")
            # 延迟扫描，确保文件写入完成
            threading.Timer(2.0, self._scan_file_delayed, [event.src_path]).start()
    
    def on_modified(self, event):
        """文件修改事件"""
        if not event.is_directory:
            self.logger.info(f"检测到文件修改: {event.src_path}")
            threading.Timer(2.0, self._scan_file_delayed, [event.src_path]).start()
    
    def _scan_file_delayed(self, file_path: str):
        """延迟扫描文件"""
        try:
            if os.path.exists(file_path):
                result = self.file_filter.scan_file(file_path)
                if result.get('risk_level', 0) >= 2:
                    self.logger.warning(f"检测到风险文件: {file_path}, 风险级别: {result.get('risk_level')}")
        except Exception as e:
            self.logger.error(f"延迟扫描文件错误: {e}")